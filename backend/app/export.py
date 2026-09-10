"""Server-side Excel export — mirrors the sheets the old dashboards produced."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics.engine import (
    research_categories,
    research_matches,
    research_text,
    theme_counts,
)

HEAD_FILL = PatternFill("solid", fgColor="0F172A")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)


def _write_sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_workbook(responses: list[dict], respondent_type: str) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    if respondent_type == "faculty":
        _faculty_sheets(wb, responses)
    else:
        _student_sheets(wb, responses)

    _theme_sheet(wb, responses)
    _group_sheet(wb, responses, "department", "Department")
    if respondent_type == "student":
        _group_sheet(wb, responses, "level", "Level of Study")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _faculty_sheets(wb: Workbook, responses: list[dict]) -> None:
    rows = []
    for i, r in enumerate(responses, start=1):
        engagements = r.get("engagements") or [{}]
        for j, e in enumerate(engagements):
            rows.append(
                [
                    i if j == 0 else "",
                    r.get("name", ""),
                    r.get("email", "") or "",
                    r.get("department", ""),
                    r.get("location", ""),
                    f"Academic Engagement {j + 1:02d}" if any(e.values()) else "",
                    e.get("programme", ""),
                    e.get("semester", ""),
                    e.get("course_name", ""),
                    r.get("research_expertise", ""),
                    r.get("research_interests", ""),
                    r.get("research_focus", ""),
                    r.get("research_impact", ""),
                    r.get("vision_self", ""),
                    r.get("vision_india", ""),
                    r.get("meeting_date", ""),
                    _ts(r),
                ]
            )
    _write_sheet(
        wb.create_sheet("All Responses"),
        [
            "S.No.", "Name", "Email", "Department", "Location", "Academic Engagement",
            "Programme", "Semester", "Course Name",
            "Research Expertise / Area of Specialisation", "Current Research Interests",
            "Your Research in Focus", "From Research to Impact",
            "My Vision for My Future", "My Vision for India's Future",
            "Date of the meeting", "Submitted",
        ],
        rows,
        [8, 25, 28, 42, 15, 22, 28, 15, 35, 45, 45, 55, 55, 65, 65, 18, 22],
    )

    total_research = sum(1 for r in responses if research_text(r).strip())
    counts: dict[str, int] = {}
    for r in responses:
        for cat in research_matches(research_text(r)):
            counts[cat] = counts.get(cat, 0) + 1
    research_rows = [
        [
            name,
            counts.get(name, 0),
            f"{counts.get(name, 0) / total_research * 100:.1f}%" if total_research else "0%",
            cfg["note"],
        ]
        for name, cfg in research_categories().items()
    ]
    _write_sheet(
        wb.create_sheet("Research Landscape"),
        ["Research Domain", "Faculty Profiles", "% of Research Profiles", "Description"],
        research_rows,
        [42, 18, 22, 85],
    )


def _student_sheets(wb: Workbook, responses: list[dict]) -> None:
    rows = [
        [
            i,
            r.get("name", ""),
            r.get("email", "") or "",
            r.get("department", ""),
            r.get("location", ""),
            r.get("level", ""),
            r.get("programme", ""),
            r.get("year", ""),
            r.get("vision_self", ""),
            r.get("vision_india", ""),
            _ts(r),
        ]
        for i, r in enumerate(responses, start=1)
    ]
    _write_sheet(
        wb.create_sheet("All Responses"),
        [
            "S.No.", "Name", "Email", "Department", "Location", "Level of Study",
            "Programme", "Year / Semester",
            "My Vision for My Future", "My Vision for India's Future", "Submitted",
        ],
        rows,
        [8, 25, 28, 42, 15, 18, 32, 24, 65, 65, 22],
    )


def _theme_sheet(wb: Workbook, responses: list[dict]) -> None:
    self_themes = theme_counts([r.get("vision_self", "") for r in responses], "self")
    india_themes = theme_counts([r.get("vision_india", "") for r in responses], "india")
    rows = [["My Vision for My Future", name, count] for name, count in self_themes]
    rows += [["My Vision for India's Future", name, count] for name, count in india_themes]
    _write_sheet(
        wb.create_sheet("Theme Analysis"),
        ["Question", "Theme", "Responses"],
        rows,
        [38, 55, 14],
    )


def _group_sheet(wb: Workbook, responses: list[dict], field: str, label: str) -> None:
    groups: dict[str, list[dict]] = {}
    for r in responses:
        key = str(r.get(field) or "Unspecified").strip() or "Unspecified"
        groups.setdefault(key, []).append(r)

    rows = []
    for name in sorted(groups):
        rows_in = groups[name]
        self_t = theme_counts([r.get("vision_self", "") for r in rows_in], "self")
        india_t = theme_counts([r.get("vision_india", "") for r in rows_in], "india")
        rows.append(
            [
                name,
                len(rows_in),
                self_t[0][0] if self_t else "",
                self_t[0][1] if self_t else 0,
                "; ".join(f"{n} ({c})" for n, c in self_t[:5]),
                india_t[0][0] if india_t else "",
                india_t[0][1] if india_t else 0,
                "; ".join(f"{n} ({c})" for n, c in india_t[:5]),
            ]
        )
    _write_sheet(
        wb.create_sheet(f"{label} Insights"[:31]),
        [
            label, "Total Responses",
            "Future Vision — Leading Theme", "Count", "Future Vision — Top 5",
            "India Vision — Leading Theme", "Count", "India Vision — Top 5",
        ],
        rows,
        [45, 16, 38, 10, 75, 40, 10, 75],
    )


def _ts(r: dict) -> str:
    value = r.get("created_at")
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M")
    return str(value or "")
